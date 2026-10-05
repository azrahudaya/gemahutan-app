"""Render isi proposal (daftar blok) ke .docx dan ke HTML untuk pratinjau PDF.

Blok:
  ("h1", teks) ("h2", teks) ("p", teks) ("p0", teks tanpa inden)
  ("bul", [butir]) ("num", [butir])
  ("tbl", judul, header, baris, lebar_cm)
  ("fig", path, lebar_cm, keterangan)
  ("eq", rumus, nomor) persamaan bernomor rata kanan
  ("kode", [baris]) ("refs", [pustaka]) ("pb",) ("cover", dict)
Markup sebaris: *miring*, **tebal**, [PERLU DIISI: ...] (disorot kuning).
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

POLA = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[PERLU DIISI:[^\]]*\])")
FONT = "Times New Roman"


def potong(teks):
    for bag in POLA.split(teks):
        if not bag:
            continue
        if bag.startswith("**"):
            yield bag[2:-2], "b"
        elif bag.startswith("*"):
            yield bag[1:-1], "i"
        elif bag.startswith("[PERLU DIISI"):
            yield bag, "h"
        else:
            yield bag, ""


# ------------------------------------------------------------------ DOCX
def _font(run, ukuran=None, tebal=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if ukuran:
        run.font.size = Pt(ukuran)
    if tebal is not None:
        run.bold = tebal


def _isi(par, teks, ukuran=None, tebal=None):
    for t, gaya in potong(teks):
        r = par.add_run(t)
        _font(r, ukuran, tebal)
        if gaya == "b":
            r.bold = True
        elif gaya == "i":
            r.italic = True
        elif gaya == "h":
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return par


def _arsir(sel, warna="E4E4E7"):
    tcPr = sel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), warna)
    tcPr.append(shd)


def _nomor_halaman(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    _font(r, 11)
    for jenis, teks in (("begin", None), (None, "PAGE"), ("end", None)):
        if jenis:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), jenis)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = teks
        r._r.append(el)


def ke_docx(blok, path: Path, akar: Path):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(s, m, Cm(2.5))
    s.different_first_page_header_footer = True
    _nomor_halaman(s)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing, pf.space_after, pf.space_before = 1.5, Pt(3), Pt(0)
    for nama, before in (("Heading 1", 12), ("Heading 2", 8)):
        st = doc.styles[nama]
        st.font.name, st.font.size, st.font.bold = FONT, Pt(12), True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.italic = False
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.element.rPr.rFonts.set(qn("w:ascii"), FONT)
        st.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(4)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.keep_with_next = True

    nomor_gambar = nomor_tabel = 0
    for b in blok:
        jenis = b[0]
        if jenis == "cover":
            c = b[1]
            for teks, uk, tebal, spasi in c["baris"]:
                if teks == "LOGO":
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(akar / c["logo"]), width=Cm(4.2))
                    continue
                if teks.startswith("IMG:"):
                    rel, lebar = teks[4:].split("|")
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(akar / rel), width=Cm(float(lebar)))
                    continue
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(spasi)
                _isi(p, teks, uk, tebal)
        elif jenis in ("h1", "h2"):
            doc.add_heading(level=1 if jenis == "h1" else 2).add_run(b[1])
            h = doc.paragraphs[-1]
            for r in h.runs:
                _font(r, 12, True)
                r.font.color.rgb = RGBColor(0, 0, 0)
        elif jenis in ("p", "p0"):
            p = _isi(doc.add_paragraph(), b[1])
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if jenis == "p":
                p.paragraph_format.first_line_indent = Cm(1)
        elif jenis == "bul":
            for butir in b[1]:
                p = _isi(doc.add_paragraph(style="List Bullet"), butir)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_after = Pt(1)
        elif jenis == "num":
            for i, butir in enumerate(b[1], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.75)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
                p.paragraph_format.space_after = Pt(1)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                _isi(p, f"{i}.\t{butir}")
        elif jenis == "eq":
            from docx.enum.text import WD_TAB_ALIGNMENT
            p = doc.add_paragraph()
            tab = p.paragraph_format.tab_stops
            tab.add_tab_stop(Cm(8), WD_TAB_ALIGNMENT.CENTER)
            tab.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT)
            _isi(p, f"\t{b[1]}\t({b[2]})")
        elif jenis == "kode":
            for baris in b[1]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                p.paragraph_format.line_spacing, p.paragraph_format.space_after = 1.0, Pt(0)
                r = p.add_run(baris)
                r.font.name, r.font.size = "Courier New", Pt(10)
        elif jenis == "tbl":
            _, judul, header, baris, lebar = b
            nomor_tabel += 1
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            cap.paragraph_format.keep_with_next = True
            cap.paragraph_format.space_before = Pt(6)
            _isi(cap, f"Tabel {nomor_tabel}. {judul}", 11)
            t = doc.add_table(rows=1 + len(baris), cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_i, isi_baris in enumerate([header] + baris):
                for c_i, teks in enumerate(isi_baris):
                    sel = t.cell(r_i, c_i)
                    sel.width = Cm(lebar[c_i])
                    p = sel.paragraphs[0]
                    p.paragraph_format.line_spacing, p.paragraph_format.space_after = 1.0, Pt(0)
                    p.paragraph_format.first_line_indent = Cm(0)
                    _isi(p, teks, 10.5, True if r_i == 0 else None)
                    if r_i == 0:
                        _arsir(sel)
                    if len(baris) <= 15 and r_i < len(baris):
                        p.paragraph_format.keep_with_next = True
                trPr = t.rows[r_i]._tr.get_or_add_trPr()
                cs = OxmlElement("w:cantSplit")
                trPr.append(cs)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif jenis == "fig":
            _, rel, lebar, ket = b
            nomor_gambar += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.line_spacing = 1.0
            p.add_run().add_picture(str(akar / rel), width=Cm(lebar))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(8)
            _isi(cap, f"Gambar {nomor_gambar}. {ket}", 11)
        elif jenis == "refs":
            for ref in b[1]:
                p = _isi(doc.add_paragraph(), ref)
                p.paragraph_format.left_indent = Cm(1)
                p.paragraph_format.first_line_indent = Cm(-1)
                p.paragraph_format.line_spacing, p.paragraph_format.space_after = 1.15, Pt(4)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif jenis == "pb":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    # Templat bawaan python-docx menulis <w:zoom> tanpa atribut percent (tidak lolos skema OOXML).
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    doc.save(path)
    return nomor_gambar, nomor_tabel
