"""Instrumen validasi pakar GemaHutan (siap cetak).

    python3 docs/assets/proposal_src/buat_instrumen.py  ->  docs/instrumen-validasi-pakar.docx
Skala SUS mengikuti Brooke (1996): 10 butir, skor 1-5, hasil 0-100.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

AKAR = Path(__file__).resolve().parents[3]
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(s, m, Cm(2.2))
st = doc.styles["Normal"]
st.font.name, st.font.size = "Times New Roman", Pt(11)
st.paragraph_format.space_after, st.paragraph_format.line_spacing = Pt(3), 1.15


def judul(teks, uk=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run(teks)
    r.bold, r.font.size = True, Pt(uk)
    return p


def par(teks, miring=False):
    p = doc.add_paragraph()
    r = p.add_run(teks)
    r.italic = miring
    return p


def tabel(header, baris, lebar, tinggi=None):
    t = doc.add_table(rows=1 + len(baris), cols=len(header))
    t.style = "Table Grid"
    for i, isi in enumerate([header] + baris):
        for j, teks in enumerate(isi):
            c = t.cell(i, j)
            c.width = Cm(lebar[j])
            run = c.paragraphs[0].add_run(teks)
            run.font.size = Pt(10)
            run.bold = i == 0
        if tinggi and i > 0:
            t.rows[i].height = Cm(tinggi)
    doc.add_paragraph()
    return t


p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("INSTRUMEN VALIDASI PAKAR\nGemaHutan: Platform Pemantauan Akustik Pasif")
r.bold, r.font.size = True, Pt(13)
par("Durasi sekitar 45 menit: wawancara 15 menit, mencoba aplikasi 15 menit, menilai deteksi dan kuesioner "
    "15 menit. Minta izin sebelum merekam atau mengutip nama pakar.", miring=True)

judul("A. Identitas Pakar")
tabel(["Butir", "Isian"], [[x, ""] for x in (
    "Nama lengkap dan gelar", "Instansi", "Bidang keahlian (ornitologi, bioakustik, pengelolaan kawasan)",
    "Pengalaman di bidang ini (tahun)", "Tanggal dan cara sesi (tatap muka / daring)",
    "Boleh dikutip namanya di proposal? (ya / tidak / hanya instansi)")], [7.5, 9.1], 0.8)

judul("B. Wawancara Kebutuhan")
for i, q in enumerate([
    "Bagaimana pemantauan burung dilakukan saat ini di tempat Anda, dan apa kendala terbesarnya?",
    "Berapa banyak data (jam rekaman atau hari survei) per bulan, dan siapa yang menganalisisnya?",
    "Apakah 25 spesies target GemaHutan relevan? Spesies apa yang paling perlu ditambahkan dan mengapa?",
    "Siapa yang biasanya memverifikasi identifikasi spesies, dan berapa lama waktu yang dibutuhkan?",
    "Informasi apa yang paling dibutuhkan pengelola kawasan: daftar spesies, tren waktu, peta, atau yang lain?",
    "Seberapa besar kesalahan deteksi otomatis yang masih dapat diterima, asalkan ada verifikasi ahli?",
    "Apa risiko terbesar jika sistem seperti ini dipakai di lapangan?",
], 1):
    par(f"{i}. {q}")
    par("Jawaban: ....................................................................................................."
        ".................................................................................................")

judul("C. Mencoba Aplikasi (catat tanpa membantu)")
tabel(["Tugas", "Berhasil (ya/tidak)", "Waktu (menit)", "Catatan kesulitan"], [
    ["1. Masuk dan buka Dashboard", "", "", ""],
    ["2. Unggah satu rekaman ke sebuah lokasi", "", "", ""],
    ["3. Buka detail rekaman, putar audio, baca spektrogram", "", "", ""],
    ["4. Verifikasi 10 deteksi (benar / salah / ragu)", "", "", ""],
    ["5. Unduh CSV dan buka halaman Laporan", "", "", ""],
], [6.6, 2.8, 2.4, 4.8], 0.9)

judul("D. Penilaian Kebenaran Deteksi")
par("Pilih 20 deteksi dari rekaman nyata (utamakan rekaman lapangan tim). Pakar mendengarkan potongan audionya, "
    "lalu mengisi penilaian. Hasilnya menjadi angka presisi yang dinilai ahli untuk Bagian 7.3 proposal.")
tabel(["No", "Rekaman / waktu", "Spesies prediksi", "Keyakinan", "Penilaian pakar (B/S/R)", "Spesies sebenarnya"],
      [[str(i), "", "", "", "", ""] for i in range(1, 21)], [1.0, 3.4, 3.8, 2.0, 3.0, 3.4], 0.6)

judul("E. Penilaian Umum (1 = sangat buruk, 5 = sangat baik)")
tabel(["Aspek", "1", "2", "3", "4", "5"], [[a, "", "", "", "", ""] for a in (
    "Kesesuaian dengan kebutuhan pemantauan", "Kemudahan penggunaan", "Kejelasan informasi hasil",
    "Kepercayaan terhadap hasil deteksi (setelah verifikasi)", "Kesiapan dipakai di lapangan",
    "Manfaat bagi konservasi")], [8.6, 1.6, 1.6, 1.6, 1.6, 1.6], 0.7)

judul("F. System Usability Scale (SUS)")
par("1 = sangat tidak setuju, 5 = sangat setuju.")
sus = [
    "Saya pikir saya akan sering menggunakan sistem ini.",
    "Saya merasa sistem ini terlalu rumit.",
    "Saya merasa sistem ini mudah digunakan.",
    "Saya merasa membutuhkan bantuan orang teknis untuk dapat menggunakan sistem ini.",
    "Saya merasa berbagai fungsi dalam sistem ini terpadu dengan baik.",
    "Saya merasa terlalu banyak hal yang tidak konsisten dalam sistem ini.",
    "Saya membayangkan kebanyakan orang akan cepat belajar menggunakan sistem ini.",
    "Saya merasa sistem ini sangat merepotkan untuk digunakan.",
    "Saya merasa sangat percaya diri menggunakan sistem ini.",
    "Saya perlu mempelajari banyak hal sebelum dapat menggunakan sistem ini.",
]
tabel(["No", "Pernyataan", "1", "2", "3", "4", "5"], [[str(i), t, "", "", "", "", ""] for i, t in enumerate(sus, 1)],
      [1.0, 10.6, 1.0, 1.0, 1.0, 1.0, 1.0], 0.7)
par("Cara menghitung: butir ganjil = skor - 1; butir genap = 5 - skor; jumlahkan lalu kalikan 2,5 (hasil 0-100). "
    "Skor di atas 68 umumnya dianggap di atas rata-rata.")

judul("G. Masukan Terbuka")
par("Hal yang paling perlu diperbaiki, fitur yang paling berguna, dan kalimat yang boleh dikutip di proposal:")
for _ in range(4):
    par("....................................................................................................."
        "..................................................................................................")

judul("H. Cara Memakai Hasilnya di Proposal")
for t in ("Kutipan kebutuhan (B1, B5) memperkuat Latar Belakang dan aspek Urgensi.",
          "Presisi hasil penilaian pakar (D) ditambahkan ke Tabel 5 sebagai 'presisi dinilai pakar'.",
          "Skor SUS (F) dan hasil tugas (C) ditambahkan ke Bagian 6.3 atau 8 sebagai bukti kegunaan.",
          "Nama dan instansi pakar mengisi penanda [PERLU DIISI] di Bagian 5, hanya jika pakar mengizinkan."):
    par("• " + t)

doc.settings.element.find(qn("w:zoom")).set(qn("w:percent"), "100")
out = AKAR / "docs/instrumen-validasi-pakar.docx"
doc.save(out)
print(out)
