"""Surat pernyataan ketua pengusul (Lampiran 2 panduan USB 2026) sebagai berkas terpisah.

    python3 docs/assets/proposal_src/buat_surat.py  ->  docs/surat-pernyataan-usb.docx
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from isi_1 import JUDUL

AKAR = Path(__file__).resolve().parents[3]
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(s, m, Cm(2.5))
st = doc.styles["Normal"]
st.font.name, st.font.size = "Times New Roman", Pt(12)
st.paragraph_format.line_spacing, st.paragraph_format.space_after = 1.5, Pt(0)


def par(teks="", tebal=False, rata=WD_ALIGN_PARAGRAPH.JUSTIFY, sorot=False):
    p = doc.add_paragraph()
    p.alignment = rata
    r = p.add_run(teks)
    r.bold = tebal
    if sorot:
        from docx.enum.text import WD_COLOR_INDEX
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


par("SURAT PERNYATAAN KETUA PENGUSUL", True, WD_ALIGN_PARAGRAPH.CENTER)
par()
par("Yang bertanda tangan di bawah ini:")
for label, isi in (("Nama", "Artya Epriliani"), ("NPM", "10124205"), ("Universitas", "Gunadarma"),
                   ("Fakultas", "Ilmu Komputer dan Teknologi Informasi"), ("Program Studi", "Sistem Informasi")):
    p = par(f"{label}\t: {isi}")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(4))
par()
par(f"dengan ini menyatakan bahwa proposal saya dengan judul \"{JUDUL}\" yang diusulkan dalam lomba USB "
    "(Unlimited Software Building) 2026 bersifat orisinal dan belum pernah diikutkan pada lomba lain, serta "
    "naskahnya belum pernah dipublikasikan dalam bentuk prosiding atau jurnal.")
par()
par("Bilamana di kemudian hari ditemukan ketidaksesuaian dengan pernyataan ini, saya bersedia dituntut dan diproses "
    "sesuai dengan ketentuan yang berlaku serta mengembalikan seluruh hadiah yang sudah diterima kepada Universitas "
    "Gunadarma.")
par()
par("Demikian pernyataan ini dibuat dengan sesungguhnya dan dengan sebenar-benarnya.")
par()
t = doc.add_table(rows=6, cols=2)
isi = [("", "Jakarta, [PERLU DIISI: tanggal]"), ("Mengetahui,", "Yang menyatakan,"), ("Ketua USB", "Materai Rp10.000"),
       ("", ""), ("(Dr. Witta Listiya Ningrum, S.Kom., MMSI)", "(Artya Epriliani)"), ("", "NPM 10124205")]
for i, (a, b) in enumerate(isi):
    for j, teks in enumerate((a, b)):
        p = t.cell(i, j).paragraphs[0]
        r = p.add_run(teks)
        if "PERLU DIISI" in teks:
            from docx.enum.text import WD_COLOR_INDEX
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if i == 3:
        t.rows[i].height = Cm(2.2)
doc.settings.element.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}zoom").set(
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}percent", "100")
out = AKAR / "docs/surat-pernyataan-usb.docx"
doc.save(out)
print(out)
